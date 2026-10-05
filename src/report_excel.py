"""Build the Excel operations report.

Derived columns (trailers, capacity status, lookups, KPI status) are live Excel formulas
that read from the data sheets and the named settings, so the workbook recalculates if
a planner edits a forecast, a pallet max, or a target.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import config

HEADER_FILL = PatternFill("solid", fgColor="184F95")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color="0B0B0B")
SUBTITLE_FONT = Font(italic=True, size=10, color="52514E")
SECTION_FONT = Font(bold=True, size=12, color="184F95")
TILE_LABEL_FONT = Font(size=9, color="52514E")
TILE_VALUE_FONT = Font(bold=True, size=16, color="0B0B0B")
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")
TILE_FILL = PatternFill("solid", fgColor="EEF4FC")
RED_FILL = PatternFill("solid", fgColor="F8D7D7")
RED_FONT = Font(color="9C1C1C")
GREEN_FILL = PatternFill("solid", fgColor="DAF2DA")
GREEN_FONT = Font(color="1E6B1E")
AMBER_FILL = PatternFill("solid", fgColor="FDEBC8")
AMBER_FONT = Font(color="8A5300")
THIN = Side(style="thin", color="C3C2B7")

PCT, PCT1, MONEY, MONEY0, INT, DEC1, DATE = "0%", "0.0%", "$#,##0.00", "$#,##0", "#,##0", "#,##0.0", "mmm d, yyyy"
FC = "'Volume Forecast'"


def _title(ws, title, subtitle=None):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = SUBTITLE_FONT


def _header(ws, row, col, headers):
    for i, h in enumerate(headers):
        cell = ws.cell(row=row, column=col + i, value=h)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _widths(ws, widths):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def _formats(ws, first_row, last_row, formats):
    for col, fmt in formats.items():
        for r in range(first_row, last_row + 1):
            ws[f"{col}{r}"].number_format = fmt


def _clean(value):
    """Convert pandas values to something openpyxl can write."""
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, str):
        return value
    if pd.isna(value):
        return None
    return value.item() if hasattr(value, "item") else value


def _pretty(name: str) -> str:
    """Turn a column name like 'trailer_fill_pct' into 'Trailer fill %'."""
    if " " in name:
        return name
    words = name.replace("_pct", " %").replace("_id", " ID").replace("_", " ")
    return words[:1].upper() + words[1:]


def _write_frame(ws, frame, start_row=4, start_col=1, table_name=None, formats=None, pretty=True):
    """Write a DataFrame with styled headers; optionally register it as an Excel Table.

    Data tables that PivotTables read keep their raw column names (pretty=False).
    """
    _header(ws, start_row, start_col, [_pretty(c) if pretty else c for c in frame.columns])
    for r, row in enumerate(frame.itertuples(index=False), start=start_row + 1):
        for c, value in enumerate(row, start=start_col):
            ws.cell(row=r, column=c, value=_clean(value))
    last_row = start_row + len(frame)
    last_col = start_col + len(frame.columns) - 1
    if formats:
        letters = {name: get_column_letter(start_col + i) for i, name in enumerate(frame.columns)}
        _formats(ws, start_row + 1, last_row, {letters[k]: v for k, v in formats.items() if k in letters})
    if table_name:
        ref = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(last_col)}{last_row}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(table)
    return last_row


def _status_rules(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Over capacity"'], fill=RED_FILL, font=RED_FONT))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"At risk"'], fill=AMBER_FILL, font=AMBER_FONT))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Covered"'], fill=GREEN_FILL, font=GREEN_FONT))


def _show_axes(chart):
    """Keep both axes visible and treat weekly dates as categories, not a daily date axis."""
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.legend.position = "t"
    chart.legend.overlay = False
    chart.x_axis.auto = False
    chart.x_axis.tickLblSkip = 4
    chart.x_axis.number_format = "mmm d"


def _name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


# ---------------------------------------------------------------- sheets
def _about(wb, ctx):
    ws = wb.active
    ws.title = "About"
    _title(ws, "Pool Point Network Planner", "Outbound transportation planning workbook. Built by Nana Afari.")
    lines = [
        ("What this is", "A weekly planning pack for an outbound network of one DC, 8 pool points, 4 linehaul carriers, and 136 stores."),
        ("Demand", "Seasonality and trend come from real U.S. Census clothing store sales (FRED series MRTSSM448USN)."),
        ("Simulated", "Stores, carriers, rates, and operations are simulated. Pool point cities and road miles are approximate."),
        ("As-of date", f"History runs through the week ending {ctx['as_of']:%b %d, %Y}. Forecast covers the next 26 weeks."),
        ("Forecast", f"{ctx['champion']} won a two-window backtest with {ctx['wape']:.1%} WAPE at the pool point level."),
        ("", ""),
        ("Sheet", "Use it to"),
        ("Dashboard", "Pick a pool point to see its KPIs, 26-week forecast, and trailer needs vs. contract."),
        ("KPI Scorecard", "Benchmark pool points and carriers against the targets on the Settings sheet."),
        ("Volume Forecast", "Weekly forecast by lane. Trailers and capacity status are formulas."),
        ("Forecast Matrix", "Pool point by week heatmap of forecast pallets."),
        ("Carrier Outlook", "Surge trailers to request from each carrier, plus the outlook notes to send them."),
        ("Peak Week Schedule", "Store delivery plan for the peak week after change requests and mitigation."),
        ("Store Visibility", "Every store: carrier, schedule, pallet max, next 4 weeks of volume, peak week status."),
        ("Schedule Changes", "Change requests, the impact of each one, and the scenario comparison."),
        ("Backlog", "Stores that would carry backlog in the peak week and the fix applied."),
        ("Weekly History", "208 weeks of lane operations as an Excel Table, ready for pivots."),
        ("Settings", "Targets and constants used by every formula. Yellow cells are inputs."),
    ]
    for i, (a, b) in enumerate(lines, start=4):
        ws.cell(row=i, column=1, value=a).font = Font(bold=True)
        ws.cell(row=i, column=2, value=b)
    ws["A10"].fill = ws["B10"].fill = HEADER_FILL
    ws["A10"].font = ws["B10"].font = HEADER_FONT
    _widths(ws, {"A": 22, "B": 120})


def _settings(wb, ctx):
    ws = wb.create_sheet("Settings")
    _title(ws, "Settings", "Named inputs used by formulas across the workbook. Change a yellow cell and everything recalculates.")
    rows = [
        ("As-of week ending", ctx["as_of"].to_pydatetime(), DATE, "AsOfDate"),
        ("Peak plan week ending", ctx["plan_week"].to_pydatetime(), DATE, "PlanWeek"),
        ("Pallets per trailer", config.PALLETS_PER_TRAILER, INT, "PalletsPerTrailer"),
        ("On-time delivery target", config.TARGETS["on_time_pct"], PCT, "TargetOnTime"),
        ("Trailer fill target", config.TARGETS["trailer_fill_pct"], PCT, "TargetFill"),
        ("Cost per pallet alert (x network)", 1.10, "0.00", "CostIndexAlert"),
        ("Forecast method", ctx["champion"], "@", None),
        ("Backtest WAPE (pool point)", ctx["wape"], PCT1, None),
    ]
    _header(ws, 4, 1, ["Setting", "Value"])
    for i, (label, value, fmt, name) in enumerate(rows, start=5):
        ws.cell(row=i, column=1, value=label)
        cell = ws.cell(row=i, column=2, value=value)
        cell.number_format = fmt
        if name:
            cell.fill = INPUT_FILL
            _name(wb, name, f"Settings!$B${i}")
    _widths(ws, {"A": 34, "B": 22})


def _carriers(wb, carriers):
    ws = wb.create_sheet("Carriers")
    _title(ws, "Carriers", "Linehaul carriers and contract rates (fictional).")
    _write_frame(ws, carriers, table_name="tblCarriers",
                 formats={"fixed_charge": MONEY0, "rate_per_mile": MONEY, "base_on_time": PCT1})
    _widths(ws, {c: 16 for c in "ABCDE"})


def _pool_points(wb, pool_points):
    ws = wb.create_sheet("Pool Points")
    _title(ws, "Pool Points", "Lane master data. Linehaul cost per trailer is a formula: carrier fixed charge + miles x rate.")
    cols = ["pool_point_id", "city", "state", "miles_from_dc", "carrier_id", "handling_per_pallet", "cost_per_stop",
            "linehaul_cost_per_trailer", "contracted_trailers_per_week", "trailing_fill_pct", "winter_exposed"]
    frame = pool_points[cols].copy()
    frame["winter_exposed"] = frame["winter_exposed"].map({True: "Yes", False: "No"})
    last = _write_frame(ws, frame, table_name="tblPoolPoints",
                        formats={"handling_per_pallet": MONEY, "cost_per_stop": MONEY0, "linehaul_cost_per_trailer": MONEY,
                                 "trailing_fill_pct": PCT1})
    for r in range(5, last + 1):
        ws[f"H{r}"] = (f"=_xlfn.XLOOKUP(E{r},Carriers!$A:$A,Carriers!$C:$C)"
                       f"+D{r}*_xlfn.XLOOKUP(E{r},Carriers!$A:$A,Carriers!$D:$D)")
        ws[f"H{r}"].number_format = MONEY
    _widths(ws, {c: 15 for c in "ABCDEFGHIJK"})
    return last


def _stores(wb, stores):
    ws = wb.create_sheet("Stores")
    _title(ws, "Stores", "Store master data, including stores opening in the forecast window.")
    cols = ["store_id", "pool_point_id", "market", "size_tier", "base_weekly_pallets", "delivery_days",
            "deliveries_per_week", "pallet_max", "open_date"]
    _write_frame(ws, stores[cols], table_name="tblStores", formats={"base_weekly_pallets": DEC1, "open_date": DATE})
    _widths(ws, {c: 14 for c in "ABCDEFGHI"})


def _volume_forecast(wb, pool_forecast, pool_points):
    ws = wb.create_sheet("Volume Forecast")
    _title(ws, "Volume Forecast by Lane",
           "Columns I to N are formulas: trailers = pallets / (pallets per trailer x lane fill rate), rounded up.")
    frame = pool_forecast.merge(pool_points[["pool_point_id", "city", "carrier_id"]], on="pool_point_id")
    frame = frame[["week_ending", "pool_point_id", "city", "carrier_id", "forecast_pallets", "low_pallets",
                   "high_pallets", "new_store_pallets"]].sort_values(["week_ending", "pool_point_id"])
    headers = ["Week ending", "Pool point", "City", "Carrier", "Forecast pallets", "Low (P10)", "High (P80)",
               "New store pallets", "Lane fill rate", "Trailers (forecast)", "Trailers (high)",
               "Contracted trailers", "Extra trailers needed", "Capacity status"]
    _header(ws, 4, 1, headers)
    for r, row in enumerate(frame.itertuples(index=False), start=5):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=_clean(value))
        ws[f"I{r}"] = f"=_xlfn.XLOOKUP(B{r},'Pool Points'!$A:$A,'Pool Points'!$J:$J)"
        ws[f"J{r}"] = f"=ROUNDUP(E{r}/(PalletsPerTrailer*I{r}),0)"
        ws[f"K{r}"] = f"=ROUNDUP(G{r}/(PalletsPerTrailer*I{r}),0)"
        ws[f"L{r}"] = f"=_xlfn.XLOOKUP(B{r},'Pool Points'!$A:$A,'Pool Points'!$I:$I)"
        ws[f"M{r}"] = f"=MAX(0,K{r}-L{r})"
        ws[f"N{r}"] = f'=IF(J{r}>L{r},"Over capacity",IF(K{r}>L{r},"At risk","Covered"))'
    last = 4 + len(frame)
    _formats(ws, 5, last, {"A": DATE, "E": DEC1, "F": DEC1, "G": DEC1, "H": DEC1, "I": PCT1})
    _status_rules(ws, f"N5:N{last}")
    table = Table(displayName="tblForecast", ref=f"A4:N{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "C5"
    _widths(ws, {"A": 14, "B": 11, "C": 12, "D": 10, **{c: 12 for c in "EFGHIJKLM"}, "N": 15})
    return last


def _forecast_matrix(wb, pool_points, weeks, fc_last):
    ws = wb.create_sheet("Forecast Matrix")
    _title(ws, "Forecast Matrix", "Forecast pallets by pool point and week (SUMIFS over the Volume Forecast sheet).")
    ws.cell(row=4, column=1, value="Pool point")
    ws.cell(row=4, column=2, value="City")
    _header(ws, 4, 1, ["Pool point", "City"])
    for j, week in enumerate(weeks, start=3):
        cell = ws.cell(row=4, column=j, value=week.to_pydatetime())
        cell.number_format, cell.fill, cell.font = "mmm d", HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for i, pp in enumerate(pool_points.itertuples(), start=5):
        ws.cell(row=i, column=1, value=pp.pool_point_id)
        ws.cell(row=i, column=2, value=pp.city)
        for j in range(3, 3 + len(weeks)):
            col = get_column_letter(j)
            ws[f"{col}{i}"] = (f"=SUMIFS({FC}!$E$5:$E${fc_last},{FC}!$B$5:$B${fc_last},$A{i},"
                               f"{FC}!$A$5:$A${fc_last},{col}$4)")
            ws[f"{col}{i}"].number_format = INT
    total_row = 5 + len(pool_points)
    ws.cell(row=total_row, column=1, value="Network").font = Font(bold=True)
    for j in range(3, 3 + len(weeks)):
        col = get_column_letter(j)
        ws[f"{col}{total_row}"] = f"=SUM({col}5:{col}{total_row - 1})"
        ws[f"{col}{total_row}"].number_format = INT
        ws[f"{col}{total_row}"].font = Font(bold=True)
    last_col = get_column_letter(2 + len(weeks))
    ws.conditional_formatting.add(f"C5:{last_col}{total_row - 1}",
                                  ColorScaleRule(start_type="min", start_color="FCFCFB", end_type="max", end_color="3987E5"))
    ws.freeze_panes = "C5"
    _widths(ws, {"A": 11, "B": 12, **{get_column_letter(j): 8 for j in range(3, 3 + len(weeks))}})


def _carrier_outlook(wb, carrier_weeks, carriers, weeks, fc_last, notes):
    ws = wb.create_sheet("Carrier Outlook")
    _title(ws, "Carrier Outlook", "Surge trailers to request from each carrier (high case vs. contract). Formulas sum the Volume Forecast sheet.")
    summary = (carrier_weeks.groupby(["carrier_id", "carrier_name"], as_index=False)
               .agg(lanes=("lanes", "max"), contracted_trailers=("contracted_trailers", "max"),
                    avg_forecast_trailers=("trailers_forecast", "mean"), peak_trailers_high=("trailers_high", "max"),
                    surge_weeks=("extra_trailers_needed", lambda s: int((s > 0).sum())),
                    max_extra_trailers=("extra_trailers_needed", "max")))
    peak_weeks = carrier_weeks.loc[carrier_weeks.groupby("carrier_id")["trailers_high"].idxmax(), ["carrier_id", "week_ending"]]
    summary = summary.merge(peak_weeks.rename(columns={"week_ending": "peak_week"}), on="carrier_id")
    ws["A4"] = "Summary"
    ws["A4"].font = SECTION_FONT
    last = _write_frame(ws, summary, start_row=5, formats={"avg_forecast_trailers": DEC1, "peak_week": DATE})

    start = last + 3
    ws.cell(row=start - 1, column=1, value="Extra trailers needed by week (high case)").font = SECTION_FONT
    _header(ws, start, 1, ["Carrier", "Name"])
    for j, week in enumerate(weeks, start=3):
        cell = ws.cell(row=start, column=j, value=week.to_pydatetime())
        cell.number_format, cell.fill, cell.font = "mmm d", HEADER_FILL, HEADER_FONT
    for i, c in enumerate(carriers.itertuples(), start=start + 1):
        ws.cell(row=i, column=1, value=c.carrier_id)
        ws.cell(row=i, column=2, value=c.carrier_name)
        for j in range(3, 3 + len(weeks)):
            col = get_column_letter(j)
            ws[f"{col}{i}"] = (f"=SUMIFS({FC}!$M$5:$M${fc_last},{FC}!$D$5:$D${fc_last},$A{i},"
                               f"{FC}!$A$5:$A${fc_last},{col}${start})")
    end = start + len(carriers)
    last_col = get_column_letter(2 + len(weeks))
    ws.conditional_formatting.add(f"C{start + 1}:{last_col}{end}",
                                  CellIsRule(operator="greaterThan", formula=["0"], fill=RED_FILL, font=RED_FONT))

    note_row = end + 3
    ws.cell(row=note_row - 1, column=1, value="Outlook notes to send carriers").font = SECTION_FONT
    for k, line in enumerate([ln for ln in notes.splitlines() if ln.strip()][1:], start=note_row):
        text = line.lstrip("#- ").strip()
        cell = ws.cell(row=k, column=1, value=text)
        if line.startswith("##"):
            cell.font = Font(bold=True)
    _widths(ws, {"A": 12, "B": 14, "C": 9, "D": 12, "E": 12, "F": 12, "G": 12, "H": 12, "I": 12,
                 **{get_column_letter(j): 9 for j in range(10, 3 + len(weeks))}})


def _scorecard(wb, card, carrier_card):
    ws = wb.create_sheet("KPI Scorecard")
    _title(ws, "KPI Scorecard, trailing 52 weeks",
           "Status columns are formulas against the targets on the Settings sheet.")
    ws["A3"] = "Targets:"
    ws["B3"] = "=\"On-time \"&TEXT(TargetOnTime,\"0%\")&\"   Fill \"&TEXT(TargetFill,\"0%\")&\"   Cost alert \"&TEXT(CostIndexAlert,\"0.00\")&\"x network\""
    cols = ["pool_point_id", "city", "carrier_id", "avg_weekly_pallets", "yoy_volume_growth", "on_time_pct",
            "on_time_status", "trailer_fill_pct", "fill_status", "cost_per_pallet", "cost_index_vs_network",
            "cost_status", "avg_backlog_pallets", "forecast_wape"]
    headers = ["Pool point", "City", "Carrier", "Avg weekly pallets", "YoY volume growth", "On-time %",
               "On-time status", "Trailer fill %", "Fill status", "Cost per pallet", "Cost vs network",
               "Cost status", "Avg backlog pallets", "Forecast WAPE"]
    frame = card[cols].copy()
    frame.columns = headers
    last = _write_frame(ws, frame, start_row=5, table_name="tblScorecard",
                        formats={"Avg weekly pallets": INT, "YoY volume growth": PCT1, "On-time %": PCT1,
                                 "Trailer fill %": PCT1, "Cost per pallet": MONEY, "Cost vs network": "0.00",
                                 "Avg backlog pallets": DEC1, "Forecast WAPE": PCT1})
    for r in range(6, last + 1):
        ws[f"G{r}"] = f'=IF(F{r}>=TargetOnTime,"Meets","Below target")'
        ws[f"I{r}"] = f'=IF(H{r}>=TargetFill,"Meets","Below target")'
        ws[f"L{r}"] = f'=IF(K{r}>CostIndexAlert,"High cost","In range")'
    for rng in (f"G6:G{last}", f"I6:I{last}", f"L6:L{last}"):
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Meets"'], fill=GREEN_FILL, font=GREEN_FONT))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"In range"'], fill=GREEN_FILL, font=GREEN_FONT))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Below target"'], fill=RED_FILL, font=RED_FONT))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"High cost"'], fill=RED_FILL, font=RED_FONT))

    start = last + 3
    ws.cell(row=start - 1, column=1, value="Carrier scorecard").font = SECTION_FONT
    ccols = ["carrier_id", "carrier_name", "lanes", "avg_weekly_pallets", "on_time_pct", "on_time_status",
             "trailer_fill_pct", "cost_per_pallet"]
    cframe = carrier_card[ccols].copy()
    cframe.columns = ["Carrier", "Name", "Lanes", "Avg weekly pallets", "On-time %", "On-time status",
                      "Trailer fill %", "Cost per pallet"]
    clast = _write_frame(ws, cframe, start_row=start, formats={"Avg weekly pallets": INT, "On-time %": PCT1,
                                                              "Trailer fill %": PCT1, "Cost per pallet": MONEY})
    for r in range(start + 1, clast + 1):
        ws[f"F{r}"] = f'=IF(E{r}>=TargetOnTime,"Meets","Below target")'
    ws.conditional_formatting.add(f"F{start + 1}:F{clast}", CellIsRule(operator="equal", formula=['"Meets"'], fill=GREEN_FILL, font=GREEN_FONT))
    ws.conditional_formatting.add(f"F{start + 1}:F{clast}", CellIsRule(operator="equal", formula=['"Below target"'], fill=RED_FILL, font=RED_FONT))
    _widths(ws, {"A": 11, "B": 12, "C": 10, **{c: 13 for c in "DEFGHIJKLMN"}})


def _dashboard(wb, pool_points, weeks, fc_last):
    ws = wb.create_sheet("Dashboard", 1)
    _title(ws, "Pool Point Dashboard", "Pick a pool point in the yellow cell. Every number and chart below updates.")
    ws["A4"], ws["A5"], ws["A6"] = "Pool point", "City", "Carrier"
    for c in ("A4", "A5", "A6"):
        ws[c].font = Font(bold=True)
    ws["B4"] = "PP03"
    ws["B4"].fill = INPUT_FILL
    ws["B4"].font = Font(bold=True, size=12)
    dv = DataValidation(type="list", formula1=f"='Pool Points'!$A$5:$A${4 + len(pool_points)}", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B4")
    ws["B5"] = "=_xlfn.XLOOKUP($B$4,'Pool Points'!$A:$A,'Pool Points'!$B:$B)&\", \"&_xlfn.XLOOKUP($B$4,'Pool Points'!$A:$A,'Pool Points'!$C:$C)"
    ws["B6"] = "=_xlfn.XLOOKUP($B$4,'Pool Points'!$A:$A,'Pool Points'!$E:$E)"

    sc = "'KPI Scorecard'"
    tiles = [
        ("Avg weekly pallets", f"=_xlfn.XLOOKUP($B$4,{sc}!$A:$A,{sc}!$D:$D)", INT),
        ("YoY volume growth", f"=_xlfn.XLOOKUP($B$4,{sc}!$A:$A,{sc}!$E:$E)", PCT1),
        ("On-time delivery", f"=_xlfn.XLOOKUP($B$4,{sc}!$A:$A,{sc}!$F:$F)", PCT1),
        ("Trailer fill", f"=_xlfn.XLOOKUP($B$4,{sc}!$A:$A,{sc}!$H:$H)", PCT1),
        ("Cost per pallet", f"=_xlfn.XLOOKUP($B$4,{sc}!$A:$A,{sc}!$J:$J)", MONEY),
        ("Forecast WAPE", f"=_xlfn.XLOOKUP($B$4,{sc}!$A:$A,{sc}!$N:$N)", PCT1),
        ("Peak forecast week", f"=_xlfn.XLOOKUP(MAX($B$13:$B${12 + len(weeks)}),$B$13:$B${12 + len(weeks)},$A$13:$A${12 + len(weeks)})", "mmm d"),
        ("Weeks over capacity", f'=COUNTIF($G$13:$G${12 + len(weeks)},"Over capacity")', INT),
    ]
    for j, (label, formula, fmt) in enumerate(tiles, start=1):
        col = get_column_letter(j)
        ws[f"{col}8"], ws[f"{col}9"] = label, formula
        ws[f"{col}8"].font, ws[f"{col}9"].font = TILE_LABEL_FONT, TILE_VALUE_FONT
        ws[f"{col}9"].number_format = fmt
        for r in (8, 9):
            ws[f"{col}{r}"].fill = TILE_FILL
            ws[f"{col}{r}"].alignment = Alignment(horizontal="center")
    ws.conditional_formatting.add("C9", FormulaRule(formula=["C9<TargetOnTime"], font=Font(bold=True, size=16, color="9C1C1C")))
    ws.conditional_formatting.add("D9", FormulaRule(formula=["D9<TargetFill"], font=Font(bold=True, size=16, color="9C1C1C")))

    ws["A11"] = "Next 26 weeks"
    ws["A11"].font = SECTION_FONT
    _header(ws, 12, 1, ["Week ending", "Forecast pallets", "High (P80)", "Trailers needed", "Contracted",
                        "Gap (trailers)", "Capacity status"])
    for i, week in enumerate(weeks, start=13):
        ws[f"A{i}"] = week.to_pydatetime()
        ws[f"A{i}"].number_format = "mmm d, yyyy"
        crit = f"{FC}!$B$5:$B${fc_last},$B$4,{FC}!$A$5:$A${fc_last},$A{i}"
        ws[f"B{i}"] = f"=SUMIFS({FC}!$E$5:$E${fc_last},{crit})"
        ws[f"C{i}"] = f"=SUMIFS({FC}!$G$5:$G${fc_last},{crit})"
        ws[f"D{i}"] = f"=SUMIFS({FC}!$J$5:$J${fc_last},{crit})"
        ws[f"E{i}"] = "=_xlfn.XLOOKUP($B$4,'Pool Points'!$A:$A,'Pool Points'!$I:$I)"
        ws[f"F{i}"] = f"=D{i}-E{i}"
        ws[f"G{i}"] = f'=IF(D{i}>E{i},"Over capacity",IF(SUMIFS({FC}!$K$5:$K${fc_last},{crit})>E{i},"At risk","Covered"))'
        for c, fmt in zip("BCDEF", [DEC1, DEC1, INT, INT, '+#,##0;-#,##0;0']):
            ws[f"{c}{i}"].number_format = fmt
    last = 12 + len(weeks)
    _status_rules(ws, f"G13:G{last}")

    # Chart titles live in cells above each chart so they never overlap the plot
    ws["I11"] = "Forecast pallets per week"
    ws["I27"] = "Trailers needed vs. contracted per week"
    ws["I11"].font = ws["I27"].font = SECTION_FONT
    line = LineChart()
    line.height, line.width = 7.5, 17
    line.add_data(Reference(ws, min_col=2, min_row=12, max_row=last), titles_from_data=True)
    line.add_data(Reference(ws, min_col=3, min_row=12, max_row=last), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=13, max_row=last))
    line.x_axis.number_format = "mmm d"
    line.series[0].graphicalProperties.line.solidFill = "2A78D6"
    line.series[0].graphicalProperties.line.width = 28000
    line.series[1].graphicalProperties.line.solidFill = "EB6834"
    line.series[1].graphicalProperties.line.dashStyle = "dash"
    line.legend.position = "b"
    _show_axes(line)
    ws.add_chart(line, "I12")

    bar = BarChart()
    bar.height, bar.width = 7.5, 17
    bar.add_data(Reference(ws, min_col=4, min_row=12, max_row=last), titles_from_data=True)
    bar.add_data(Reference(ws, min_col=5, min_row=12, max_row=last), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=13, max_row=last))
    bar.x_axis.number_format = "mmm d"
    bar.series[0].graphicalProperties.solidFill = "2A78D6"
    bar.series[1].graphicalProperties.solidFill = "C3C2B7"
    bar.gapWidth = 40
    bar.legend.position = "b"
    _show_axes(bar)
    ws.add_chart(bar, "I28")
    _widths(ws, {c: 17 for c in "ABCDEFGH"})
    ws.freeze_panes = "A8"


def _schedule_sheet(wb, plan):
    ws = wb.create_sheet("Peak Week Schedule")
    _title(ws, f"Peak Week Delivery Schedule, week ending {plan['plan_week']:%b %d, %Y}",
           "Pallets per delivery after change requests and mitigation. Orange cells are at the store's pallet max.")
    grid = plan["deliveries"].pivot_table(index="store_id", columns="day", values="pallets", aggfunc="sum")
    grid = grid.reindex(columns=config.DAYS)
    vis = plan["store_visibility"].set_index("store_id")
    frame = vis[["pool_point_id", "delivery_days", "pallet_max"]].join(grid).reset_index()
    headers = ["Store", "Pool point", "Delivery days", "Pallet max", *config.DAYS, "Planned total", "Forecast", "Backlog"]
    _header(ws, 4, 1, headers)
    for r, row in enumerate(frame.itertuples(index=False), start=5):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=_clean(value))
        ws[f"K{r}"] = f"=SUM(E{r}:J{r})"
        ws[f"L{r}"] = vis.at[row[0], "forecast_pallets"]
        ws[f"M{r}"] = f"=MAX(0,ROUND(L{r},0)-K{r})"
        ws[f"L{r}"].number_format = DEC1
    last = 4 + len(frame)
    ws.conditional_formatting.add(f"E5:J{last}", FormulaRule(formula=["AND(E5<>\"\",E5=$D5)"], fill=AMBER_FILL, font=AMBER_FONT))
    ws.conditional_formatting.add(f"M5:M{last}", CellIsRule(operator="greaterThan", formula=["0"], fill=RED_FILL, font=RED_FONT))
    total = last + 1
    ws[f"A{total}"] = "Total"
    ws[f"A{total}"].font = Font(bold=True)
    for c in "EFGHIJKM":
        ws[f"{c}{total}"] = f"=SUM({c}5:{c}{last})"
        ws[f"{c}{total}"].font = Font(bold=True)
    ws.freeze_panes = "B5"
    _widths(ws, {"A": 10, "B": 11, "C": 15, "D": 10, **{c: 7 for c in "EFGHIJ"}, "K": 12, "L": 10, "M": 10})


def _store_visibility(wb, plan, store_forecast, stores, as_of):
    data = wb.create_sheet("Store Forecast")
    _title(data, "Store Forecast", "Weekly forecast pallets by store (pool point forecast split by each store's recent share).")
    sf_last = _write_frame(data, store_forecast, table_name="tblStoreForecast",
                           formats={"week_ending": DATE, "forecast_pallets": DEC1})
    _widths(data, {"A": 14, "B": 12, "C": 10, "D": 14})

    ws = wb.create_sheet("Store Visibility", 3)
    _title(ws, "Store Visibility Report",
           "One row per store. Carrier, next-4-week volume, and status are formulas.")
    next_weeks = pd.date_range(as_of + pd.Timedelta(weeks=1), periods=4, freq="W-SAT")
    headers = ["Store", "Pool point", "Market", "Carrier", "Size", "Delivery days", "Pallet max", "Open date",
               *[f"Wk {d:%b %d}" for d in next_weeks], "Peak wk forecast", "Peak wk planned", "Peak wk backlog", "Status"]
    _header(ws, 4, 1, headers)
    for i, d in enumerate(next_weeks):
        ws.cell(row=3, column=9 + i, value=d.to_pydatetime()).number_format = "mmm d"
    ws["H3"] = "Week ending:"
    vis = plan["store_visibility"].merge(stores[["store_id", "open_date"]], on="store_id", how="left")
    sfr = f"'Store Forecast'!$D$5:$D${sf_last}"
    for r, row in enumerate(vis.itertuples(index=False), start=5):
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = row.store_id, row.pool_point_id, row.market
        ws[f"D{r}"] = f"=_xlfn.XLOOKUP(B{r},'Pool Points'!$A:$A,'Pool Points'!$E:$E)"
        ws[f"E{r}"], ws[f"F{r}"], ws[f"G{r}"] = row.size_tier, row.delivery_days, row.pallet_max
        ws[f"H{r}"] = row.open_date.to_pydatetime() if pd.notna(row.open_date) else plan["plan_week"].to_pydatetime()
        ws[f"H{r}"].number_format = DATE
        for i in range(4):
            col = get_column_letter(9 + i)
            ws[f"{col}{r}"] = (f"=SUMIFS({sfr},'Store Forecast'!$C$5:$C${sf_last},$A{r},"
                               f"'Store Forecast'!$A$5:$A${sf_last},{col}$3)")
            ws[f"{col}{r}"].number_format = DEC1
        ws[f"M{r}"], ws[f"N{r}"], ws[f"O{r}"] = row.forecast_pallets, row.planned_pallets, row.backlog_pallets
        ws[f"M{r}"].number_format = DEC1
        ws[f"P{r}"] = (f'=IF(H{r}>AsOfDate,"Opens "&TEXT(H{r},"mmm d"),'
                       f'IF(O{r}>0,"Backlog in peak week","On plan"))')
    last = 4 + len(vis)
    ws.conditional_formatting.add(f"P5:P{last}", CellIsRule(operator="equal", formula=['"On plan"'], fill=GREEN_FILL, font=GREEN_FONT))
    ws.conditional_formatting.add(f"P5:P{last}", CellIsRule(operator="equal", formula=['"Backlog in peak week"'], fill=RED_FILL, font=RED_FONT))
    ws.conditional_formatting.add(f"P5:P{last}", FormulaRule(formula=['LEFT(P5,5)="Opens"'], fill=AMBER_FILL, font=AMBER_FONT))
    table = Table(displayName="tblStoreVisibility", ref=f"A4:P{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "B5"
    _widths(ws, {"A": 9, "B": 10, "C": 12, "D": 9, "E": 9, "F": 13, "G": 9, "H": 13,
                 **{c: 10 for c in "IJKL"}, "M": 11, "N": 11, "O": 11, "P": 20})


def _changes_sheet(wb, changes, plan):
    ws = wb.create_sheet("Schedule Changes")
    _title(ws, f"Schedule Changes, peak week ending {plan['plan_week']:%b %d, %Y}",
           "Change requests received, the impact of each one alone, and the combined scenarios.")
    ws["A4"] = "Change log"
    ws["A4"].font = SECTION_FONT
    # Long text goes in the last column of each table so it can overflow into empty cells
    log = changes[["change_id", "change_type", "pool_point_id", "store_id", "day", "old_value", "new_value",
                   "requested_by", "reason"]]
    last = _write_frame(ws, log, start_row=5)

    start = last + 3
    ws.cell(row=start - 1, column=1, value="Impact of each change on its own (vs. baseline plan)").font = SECTION_FONT
    last = _write_frame(ws, plan["per_change"], start_row=start, formats={"cost_change": '+$#,##0;-$#,##0;$0'})

    start = last + 3
    ws.cell(row=start - 1, column=1, value="Scenario comparison").font = SECTION_FONT
    last = _write_frame(ws, plan["scenarios"], start_row=start,
                        formats={"trailer_fill_pct": PCT1, "linehaul_cost": MONEY0, "total_cost": MONEY0,
                                 "cost_per_pallet": MONEY})
    _widths(ws, {"A": 24, "B": 24, **{c: 14 for c in "CDEFGIJ"}, "H": 26})


def _backlog_sheet(wb, plan):
    ws = wb.create_sheet("Backlog")
    _title(ws, "Peak Week Backlog Report",
           "Stores whose pallets exceed delivery capacity after change requests, and the mitigation applied.")
    actions = plan["mitigation_actions"].copy()
    after = plan["store_visibility"].set_index("store_id")["backlog_pallets"]
    actions["backlog_after"] = actions["store_id"].map(after).fillna(0).astype(int)
    actions = actions.sort_values("backlog_before", ascending=False)
    last = _write_frame(ws, actions, start_row=4, table_name="tblBacklog")
    ws.conditional_formatting.add(f"C5:C{last}", ColorScaleRule(start_type="min", start_color="FCFCFB", end_type="max", end_color="E34948"))
    _widths(ws, {c: 16 for c in "ABCDEFG"})


def _history(wb, pool_weeks):
    ws = wb.create_sheet("Weekly History")
    _title(ws, "Weekly Lane History", "208 weeks of simulated lane operations. Insert > PivotTable on tblWeekly to slice it.")
    frame = pool_weeks.assign(year=pool_weeks["week_ending"].dt.year, month=pool_weeks["week_ending"].dt.month)
    cols = ["week_ending", "year", "month", "pool_point_id", "carrier_id", "demand_pallets", "pallets_shipped",
            "backlog_pallets", "trailers", "trailer_fill_pct", "stops", "on_time_stops", "on_time_pct",
            "linehaul_cost", "handling_cost", "final_mile_cost", "total_cost", "cost_per_pallet"]
    _write_frame(ws, frame[cols], table_name="tblWeekly", pretty=False,
                 formats={"week_ending": DATE, "trailer_fill_pct": PCT1, "on_time_pct": PCT1, "linehaul_cost": MONEY0,
                          "handling_cost": MONEY0, "final_mile_cost": MONEY0, "total_cost": MONEY0, "cost_per_pallet": MONEY})
    ws.freeze_panes = "E5"
    _widths(ws, {get_column_letter(i): 13 for i in range(1, len(cols) + 1)})


def build_workbook(path, ctx):
    wb = Workbook()
    _about(wb, ctx)
    _settings(wb, ctx)
    _carriers(wb, ctx["carriers"])
    _pool_points(wb, ctx["pool_points"])
    _stores(wb, ctx["stores"])
    fc_last = _volume_forecast(wb, ctx["pool_forecast"], ctx["pool_points"])
    weeks = pd.DatetimeIndex(sorted(ctx["pool_forecast"]["week_ending"].unique()))
    _forecast_matrix(wb, ctx["pool_points"], weeks, fc_last)
    _carrier_outlook(wb, ctx["carrier_weeks"], ctx["carriers"], weeks, fc_last, ctx["carrier_notes"])
    _scorecard(wb, ctx["scorecard"], ctx["carrier_scorecard"])
    _dashboard(wb, ctx["pool_points"], weeks, fc_last)
    _schedule_sheet(wb, ctx["plan"])
    _store_visibility(wb, ctx["plan"], ctx["store_forecast"], ctx["stores_all"], ctx["as_of"])
    _changes_sheet(wb, ctx["changes"], ctx["plan"])
    _backlog_sheet(wb, ctx["plan"])
    _history(wb, ctx["pool_weeks"])

    order = ["About", "Dashboard", "KPI Scorecard", "Volume Forecast", "Forecast Matrix", "Carrier Outlook",
             "Store Visibility", "Peak Week Schedule", "Schedule Changes", "Backlog", "Weekly History",
             "Store Forecast", "Pool Points", "Carriers", "Stores", "Settings"]
    wb._sheets = [wb[name] for name in order]
    for name, color in [("Dashboard", "2A78D6"), ("KPI Scorecard", "2A78D6"), ("Settings", "EDA100")]:
        wb[name].sheet_properties.tabColor = color
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == "Dashboard"
    wb.active = 1
    wb.calculation.fullCalcOnLoad = True
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
